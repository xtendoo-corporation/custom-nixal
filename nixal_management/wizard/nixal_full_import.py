import base64
import calendar
import io
import re
from collections import defaultdict
from datetime import date, datetime

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl ships with Odoo's requirements.txt
    openpyxl = None

SKIP_LABELS = (
    "check", "caja", "ptes", "saldo", "acumulado", "ratio",
    "capital pendiente", "movimiento mensual", "suma saldos",
)
INCOME_HEADER_HINTS = ("venta", "ingreso", "cobro")
EXPENSE_HEADER_HINTS = ("gasto", "compra")
AGRI_SPECIES = ["Vacuno", "Porcino", "Ovino"]
MONTH_MAP = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}
YEAR_BLOCK_RE = re.compile(r"^(\d{4})(?:/(\d{4}))?$")


class NixalFullImportWizard(models.TransientModel):
    _name = "nixal.full.import.wizard"
    _description = "Importar todos los datos de los cuadros de gestión Nixal"

    file_general = fields.Binary(string="Nixal - cuadro de gestión.xlsx")
    filename_general = fields.Char()
    file_agri = fields.Binary(string="NixAl Agri - Cuadro gestión.xlsx")
    filename_agri = fields.Char()
    result_message = fields.Text(readonly=True)

    def _compute_display_name(self):
        for wizard in self:
            wizard.display_name = _("Importar todos los datos Nixal")

    # ------------------------------------------------------------------
    # Generic helpers
    # ------------------------------------------------------------------
    def _load_workbook(self, data):
        if openpyxl is None:
            raise UserError(_("El servidor no tiene instalada la librería openpyxl."))
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(data)), data_only=True)

    def _get_company(self, name):
        company = self.env["res.company"].search([("name", "=", name)], limit=1)
        if not company:
            raise UserError(_(
                "No existe la compañía '%s'. Ejecuta antes el asistente "
                "'Crear compañías Nixal'."
            ) % name)
        return company

    def _allow_company(self, company):
        """Return a copy of the wizard whose environment has `company`
        among the *allowed* companies for this request, regardless of
        which company the user currently has active in the top-right
        switcher - otherwise multi-company record rules block writes to
        that company's journal entries (e.g. when posting invoices).

        This also grants the current user access to the company itself
        (res.users.company_ids): setting allowed_company_ids alone still
        raises AccessError if the user was never added to that company
        (e.g. someone other than the user who ran 'Crear compañías Nixal'
        is running this wizard).
        """
        if company not in self.env.user.company_ids:
            self.env.user.sudo().write({"company_ids": [Command.link(company.id)]})
        allowed = set(self.env.context.get("allowed_company_ids") or self.env.companies.ids)
        allowed.add(company.id)
        return self.with_context(allowed_company_ids=list(allowed))

    def _get_partner(self, name, **extra):
        name = (name or "").strip()
        partner = self.env["res.partner"].search(
            [("name", "=", name), ("company_id", "=", False)], limit=1
        )
        if not partner:
            vals = {"name": name, "is_company": True, "company_id": False}
            vals.update(extra)
            partner = partner.create(vals)
        return partner

    def _get_product(self, name, product_type="service"):
        product = self.env["product.product"].search(
            [("name", "=", name), ("company_id", "=", False)], limit=1
        )
        if not product:
            product = product.create({
                "name": name,
                "type": product_type,
                "company_id": False,
                "invoice_policy": "order",
            })
        return product

    def _year_blocks(self, ws, header_row, max_col=None):
        """Map column -> year label ('2018' or '2020/2021') by propagating
        the sparse block-start markers on `header_row` forward."""
        mapping = {}
        current = None
        for c in range(1, (max_col or ws.max_column) + 1):
            value = ws.cell(row=header_row, column=c).value
            if isinstance(value, str) and YEAR_BLOCK_RE.match(value.strip()):
                current = value.strip()
            elif isinstance(value, int) and 1900 < value < 2100:
                current = str(value)
            mapping[c] = current
        return mapping

    def _dated_columns(self, ws, date_row=2, year_row=1, max_col=None):
        """Yield (col_index, date) for every column that can be resolved to a
        real calendar date, supporting two header layouts found in the Nixal
        sheets:
        - a real datetime directly on `date_row` (e.g. 'Ing y gtos'), or
        - a month name ('Ene', 'Feb', ...) on `date_row` combined with a
          sparse year/fiscal-year block marker ('2018' or '2020/2021') on
          `year_row` (e.g. the '*-Ingresos'/'*-Gastos' sheets).
        'Totales'/'TOTALES' aggregate columns and empty headers are skipped.
        """
        today = date.today()
        year_blocks = self._year_blocks(ws, year_row, max_col)
        for c in range(1, (max_col or ws.max_column) + 1):
            value = ws.cell(row=date_row, column=c).value
            if isinstance(value, datetime):
                if value.date() <= today:
                    yield c, value.date()
                continue
            if not isinstance(value, str):
                continue
            month = MONTH_MAP.get(value.strip().lower()[:3])
            year_label = year_blocks.get(c)
            if not month or not year_label:
                continue
            match = YEAR_BLOCK_RE.match(year_label)
            if not match:
                continue
            year_start = int(match.group(1))
            year_end = int(match.group(2)) if match.group(2) else None
            year = year_start if (year_end is None or month >= 9) else year_end
            day = calendar.monthrange(year, month)[1]
            col_date = date(year, month, day)
            # Skip columns that are budget/forecast (future), not real data.
            if col_date <= today:
                yield c, col_date

    def _is_skippable_label(self, label):
        low = (label or "").strip().lower()
        if not low:
            return True
        # "Total Fra ..." rows are actual invoice data, not subtotals like
        # "Total Ventas vacuno" / "Totales" - only skip the latter.
        if low.startswith("total") and "fra" not in low:
            return True
        return any(kw in low for kw in SKIP_LABELS)

    def _create_grouped_moves(self, company, move_type, entries, product, analytic_accounts=None):
        """entries: dict[(partner_id, date, species)] -> list of (description, amount).
        `species` is the analytic category name (e.g. 'Vacuno') or None; when
        given and found in `analytic_accounts` (dict name -> analytic account
        id), every line of that group is tagged with it.
        """
        analytic_accounts = analytic_accounts or {}
        created = 0
        for (partner_id, move_date, species), lines in entries.items():
            analytic_account_id = analytic_accounts.get(species)
            analytic_distribution = (
                {str(analytic_account_id): 100.0} if analytic_account_id else False
            )
            move_vals = {
                "move_type": move_type,
                "company_id": company.id,
                "partner_id": partner_id,
                "invoice_date": move_date,
                "invoice_line_ids": [
                    (0, 0, {
                        "product_id": product.id,
                        "name": description[:250],
                        "quantity": 1,
                        "price_unit": abs(amount),
                        "tax_ids": [(6, 0, [])],
                        "analytic_distribution": analytic_distribution,
                    })
                    for description, amount in lines
                ],
            }
            move = self.env["account.move"].with_company(company).create(move_vals)
            move.action_post()
            created += 1
        return created

    # ------------------------------------------------------------------
    # Clientes (contratos NC) -> notas en el contacto
    # ------------------------------------------------------------------
    def _import_clientes(self, wb):
        if "Clientes" not in wb.sheetnames:
            return 0
        ws = wb["Clientes"]
        count = 0
        for r in range(3, ws.max_row + 1):
            name = ws.cell(row=r, column=2).value
            amount = ws.cell(row=r, column=3).value
            if not name or not isinstance(amount, (int, float)):
                continue
            currency = ws.cell(row=r, column=4).value or ""
            duration = ws.cell(row=r, column=5).value or ""
            begins = ws.cell(row=r, column=6).value
            ends = ws.cell(row=r, column=7).value
            partner = self._get_partner(str(name), customer_rank=1)
            note = _(
                "Contrato NC (importado): %(amount)s %(currency)s/año, "
                "duración %(duration)s año(s), de %(begins)s a %(ends)s."
            ) % {
                "amount": amount,
                "currency": currency,
                "duration": duration,
                "begins": begins.date() if isinstance(begins, datetime) else begins,
                "ends": ends.date() if isinstance(ends, datetime) else ends,
            }
            partner.message_post(body=note)
            count += 1
        return count

    # ------------------------------------------------------------------
    # Hojas tipo "X-Ingresos" / "X-Gastos": fila = cliente/proveedor/concepto,
    # columnas = fechas reales (se agrupa una factura por contacto y mes).
    # ------------------------------------------------------------------
    def _import_ledger_sheet(self, wb, sheet_name, company, move_type, rank_field, product):
        if sheet_name not in wb.sheetnames:
            return 0
        ws = wb[sheet_name]
        # These sheets have month names ('Ene', 'Feb', ...) on row 3 and the
        # sparse year/fiscal-year block markers on row 2.
        columns = list(self._dated_columns(ws, date_row=3, year_row=2))
        entries = defaultdict(list)
        for r in range(1, ws.max_row + 1):
            label = ws.cell(row=r, column=2).value
            if not isinstance(label, str) or self._is_skippable_label(label):
                continue
            partner = self._get_partner(label, **{rank_field: 1})
            for c, col_date in columns:
                amount = ws.cell(row=r, column=c).value
                if not isinstance(amount, (int, float)) or not amount:
                    continue
                month_key = col_date.replace(day=1)
                entries[(partner.id, month_key, None)].append((label, amount))
        return self._create_grouped_moves(company, move_type, entries, product)

    # ------------------------------------------------------------------
    # Financiación / Inversiones (fichero general): mismo layout fila/mes que
    # las hojas -Ingresos/-Gastos, pero en una sola pestaña con importe en
    # positivo o negativo (entrada o salida de caja) en vez de dos pestañas
    # separadas. Todo va a la compañía "Nixal Financiación e Inversiones".
    # Nota: esto trata el signo del importe como ingreso/gasto financiero
    # a efectos de visibilidad, no como una imputación contable de préstamo
    # (activo/pasivo) correcta - es una simplificación deliberada.
    # ------------------------------------------------------------------
    def _import_financiacion_inversiones(self, wb, sheet_name, company, product_income, product_expense):
        if sheet_name not in wb.sheetnames:
            return 0
        ws = wb[sheet_name]
        columns = list(self._dated_columns(ws, date_row=3, year_row=2))
        entries_income = defaultdict(list)
        entries_expense = defaultdict(list)
        for r in range(1, ws.max_row + 1):
            label = ws.cell(row=r, column=2).value
            if not isinstance(label, str) or self._is_skippable_label(label):
                continue
            partner = self._get_partner(label, customer_rank=1, supplier_rank=1)
            for c, col_date in columns:
                amount = ws.cell(row=r, column=c).value
                if not isinstance(amount, (int, float)) or not amount:
                    continue
                month_key = col_date.replace(day=1)
                bucket = entries_income if amount > 0 else entries_expense
                bucket[(partner.id, month_key, None)].append((label, amount))
        created = self._create_grouped_moves(company, "out_invoice", entries_income, product_income)
        created += self._create_grouped_moves(company, "in_invoice", entries_expense, product_expense)
        return created

    # ------------------------------------------------------------------
    # Ing y gtos (Agri): secciones "Ventas ..." (ingresos) / "Gastos ..." (gastos)
    # sin contacto identificado -> se usa un contacto genérico por especie.
    # ------------------------------------------------------------------
    def _get_agri_analytic_accounts(self, company):
        """name -> analytic account id, for the species accounts created by
        the 'Crear compañías Nixal' wizard (empty dict if not set up yet)."""
        accounts = self.env["account.analytic.account"].search([
            ("company_id", "=", company.id),
            ("name", "in", AGRI_SPECIES),
        ])
        return {a.name: a.id for a in accounts}

    def _import_agri_ing_y_gtos(self, wb, company, investment_company):
        sheet_name = "Ing y gtos"
        if sheet_name not in wb.sheetnames:
            return 0
        ws = wb[sheet_name]
        columns = list(self._dated_columns(ws))
        entries_income = defaultdict(list)
        entries_expense = defaultdict(list)
        entries_investment = defaultdict(list)
        section = None  # True=income, False=expense
        species = None  # 'Vacuno' / 'Porcino' / 'Ovino' / None (shared/general)
        in_investment_section = False
        stop_processing = False
        generic_partner = self._get_partner(
            "Contraparte genérica NixAl Agri (importación masiva)", customer_rank=1, supplier_rank=1
        )
        investment_partner = self._get_partner(
            "Contraparte genérica Financiación e Inversiones (importación masiva)",
            customer_rank=1, supplier_rank=1,
        )
        for r in range(1, ws.max_row + 1):
            label = ws.cell(row=r, column=2).value
            if not isinstance(label, str) or not label.strip():
                continue
            low = label.strip().lower()
            has_numbers = any(
                isinstance(ws.cell(row=r, column=c).value, (int, float))
                for c, _d in columns
            )
            if not has_numbers:
                # "Préstamos" mixes running balances ("Saldo", "Capital
                # pendiente amortizar") with actual monthly flows under the
                # same ambiguous "Movimiento mensual" label - rather than
                # risk importing a balance as if it were a transaction, we
                # just stop reading the sheet from here on.
                if low.strip() == "préstamos" or low.strip() == "prestamos":
                    stop_processing = True
                    continue
                # The clean "Inversiones" (capex) block, however, has plain
                # monthly purchase amounts - redirect it to the dedicated
                # Financiación e Inversiones company.
                in_investment_section = low.strip() == "inversiones"
                if any(kw in low for kw in INCOME_HEADER_HINTS):
                    section = True
                elif any(kw in low for kw in EXPENSE_HEADER_HINTS):
                    section = False
                if "general" in low:
                    species = None
                for name in AGRI_SPECIES:
                    if name.lower() in low:
                        species = name
                        break
                continue
            if stop_processing or self._is_skippable_label(label) or section is None:
                continue
            if in_investment_section:
                bucket = entries_investment
                partner_id = investment_partner.id
                current_species = None
            else:
                bucket = entries_income if section else entries_expense
                partner_id = generic_partner.id
                current_species = species
            for c, col_date in columns:
                amount = ws.cell(row=r, column=c).value
                if not isinstance(amount, (int, float)) or not amount:
                    continue
                month_key = col_date.replace(day=1)
                bucket[(partner_id, month_key, current_species)].append((label, amount))
        product_income = self._get_product("Ingreso Actividad Ganadera (importado)", "service")
        product_expense = self._get_product("Gasto Actividad Ganadera (importado)", "service")
        product_investment = self._get_product("Inversión Actividad Ganadera (importado)", "service")
        analytic_accounts = self._get_agri_analytic_accounts(company)
        created = self._create_grouped_moves(
            company, "out_invoice", entries_income, product_income, analytic_accounts
        )
        created += self._create_grouped_moves(
            company, "in_invoice", entries_expense, product_expense, analytic_accounts
        )
        created += self._create_grouped_moves(
            investment_company, "in_invoice", entries_investment, product_investment
        )
        return created

    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if not self.file_general and not self.file_agri:
            raise UserError(_("Sube al menos uno de los dos ficheros Excel."))

        summary = []

        if self.file_general:
            wb = self._load_workbook(self.file_general)
            nc_company = self._get_company("Nixal NC")
            trade_company = self._get_company("Nixal Trade")
            fininv_company = self._get_company("Nixal Financiación e Inversiones")
            self = self._allow_company(nc_company)
            self = self._allow_company(trade_company)
            self = self._allow_company(fininv_company)

            n_clientes = self._import_clientes(wb)
            n_nc_in = self._import_ledger_sheet(
                wb, "NC-Ingresos", nc_company, "out_invoice", "customer_rank",
                self._get_product("Ingreso NC (importado)", "service"),
            )
            n_nc_out = self._import_ledger_sheet(
                wb, "NC-Gastos", nc_company, "in_invoice", "supplier_rank",
                self._get_product("Gasto NC (importado)", "service"),
            )
            n_trade_in = self._import_ledger_sheet(
                wb, "Trade-Ingresos", trade_company, "out_invoice", "customer_rank",
                self._get_product("Ingreso Trade (importado)", "service"),
            )
            n_trade_out = self._import_ledger_sheet(
                wb, "Trade-Gastos", trade_company, "in_invoice", "supplier_rank",
                self._get_product("Gasto Trade (importado)", "service"),
            )
            n_fin = self._import_financiacion_inversiones(
                wb, "Financiación", fininv_company,
                self._get_product("Ingreso Financiación (importado)", "service"),
                self._get_product("Gasto Financiación (importado)", "service"),
            )
            n_inv = self._import_financiacion_inversiones(
                wb, "Inversiones", fininv_company,
                self._get_product("Ingreso Inversiones (importado)", "service"),
                self._get_product("Gasto Inversiones (importado)", "service"),
            )
            summary.append(_(
                "Fichero general: %(cli)s contactos NC anotados, "
                "%(nc_in)s facturas NC, %(nc_out)s facturas de proveedor NC, "
                "%(tr_in)s facturas Trade, %(tr_out)s facturas de proveedor Trade, "
                "%(fin)s asientos de Financiación, %(inv)s asientos de Inversiones."
            ) % {
                "cli": n_clientes, "nc_in": n_nc_in, "nc_out": n_nc_out,
                "tr_in": n_trade_in, "tr_out": n_trade_out,
                "fin": n_fin, "inv": n_inv,
            })

        if self.file_agri:
            wb_agri = self._load_workbook(self.file_agri)
            agri_company = self._get_company("NixAl Agri")
            fininv_company = self._get_company("Nixal Financiación e Inversiones")
            self = self._allow_company(agri_company)
            self = self._allow_company(fininv_company)
            n_agri = self._import_agri_ing_y_gtos(wb_agri, agri_company, fininv_company)
            summary.append(_(
                "Fichero NixAl Agri: %s asientos/facturas creados en 'Ing y gtos' "
                "(incluye inversiones ganaderas en Nixal Financiación e Inversiones)."
            ) % n_agri)

        message = "\n".join(summary) or _("No se ha importado ningún dato.")
        self.result_message = message
        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Importación completa Nixal"),
                "message": message,
                "sticky": True,
            },
        }
